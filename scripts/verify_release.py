#!/usr/bin/env python3
"""Install and exercise a downloaded BoxReply wheel or sdist independently.

Run with Python 3.10+: python verify_release.py --artifact FILE --work NEW_DIR
The harness uses only the standard library. It downloads the artifact's runtime
and build dependencies, then installs from that wheelhouse with --no-index.
Application commands run under the fresh interpreter with sockets disabled.
The JSON report contains relative installation locations, never user paths.
"""
from __future__ import annotations

import argparse
import ast
import csv
import datetime as dt
from email.parser import BytesParser
import hashlib
import json
import os
from pathlib import Path
import platform
import re
import shutil
import subprocess
import sys
import tarfile
import time
import venv
import zipfile


NETWORK_GUARD = """
import socket
def network_forbidden(*args, **kwargs):
    raise RuntimeError('Release acceptance forbids network access')
socket.create_connection = network_forbidden
socket.getaddrinfo = network_forbidden
socket.socket.connect = network_forbidden
socket.socket.connect_ex = network_forbidden
"""

CLI_BOOTSTRAP = NETWORK_GUARD + """
import runpy, sys
sys.argv = ['boxreply', *sys.argv[1:]]
runpy.run_module('boxreply', run_name='__main__', alter_sys=True)
"""

IMPORT_PROBE = NETWORK_GUARD + """
import importlib.metadata, json, pathlib, sys
from urllib.parse import unquote, urlparse
import boxreply
prefix = pathlib.Path(sys.prefix).resolve()
origin = pathlib.Path(boxreply.__file__).resolve()
relative = origin.relative_to(prefix)
checkout = pathlib.Path(sys.argv[1]).resolve()
assert not origin.is_relative_to(checkout), 'Import leaked from development checkout'
distribution = importlib.metadata.distribution('boxreply')
direct = json.loads(distribution.read_text('direct_url.json'))
assert not direct.get('dir_info', {}).get('editable', False), 'Editable install'
assert urlparse(direct['url']).scheme == 'file', 'Expected local artifact installation'
archive = direct.get('archive_info', {})
digest = archive.get('hashes', {}).get('sha256')
if digest is None and archive.get('hash', '').startswith('sha256='):
    digest = archive['hash'][len('sha256='):]
assert digest == sys.argv[2], 'Installed artifact hash differs from input artifact'
assert pathlib.Path(unquote(urlparse(direct['url']).path)).name == sys.argv[3]
print(json.dumps({
    'package': distribution.metadata['Name'], 'version': distribution.version,
    'module_relative_to_venv': relative.as_posix(),
    'artifact_filename': pathlib.Path(unquote(urlparse(direct['url']).path)).name,
    'artifact_sha256': digest, 'editable': False,
    'outside_checkout': True, 'isolated_interpreter': bool(sys.flags.isolated),
    'dependencies': {d.metadata['Name']: d.version for d in importlib.metadata.distributions()},
}, sort_keys=True))
"""

EDIT_WORKBOOK = NETWORK_GUARD + """
import pathlib, sys
from openpyxl import load_workbook
source, target, mode = pathlib.Path(sys.argv[1]), pathlib.Path(sys.argv[2]), sys.argv[3]
workbook = load_workbook(source)
workbook['Instructions']['B5'] = 'Release acceptance reviewer'
workbook['Instructions']['B6'] = '2026-10-15'
for row in workbook['Review'].iter_rows(min_row=2):
    row[6].value = 'no_exception'
    row[7].value = ''
if mode == 'wrong_source':
    workbook['Review']['D2'] = 'Source changed after the frozen round'
elif mode in {'exception', 'needs_correction'}:
    workbook['Review']['G2'] = mode
    workbook['Review']['H2'] = 'Acceptance test: human review required'
elif mode != 'answer':
    raise ValueError('Unknown acceptance edit')
workbook.save(target)
workbook.close()
"""


class AcceptanceFailure(RuntimeError):
    pass


def digest(path):
    value = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def artifact_contract(path):
    """Read metadata without extracting or importing the candidate package."""
    build_requires = []
    if path.name.endswith(".whl"):
        with zipfile.ZipFile(path) as archive:
            names = [name for name in archive.namelist() if name.endswith(".dist-info/METADATA")]
            if len(names) != 1:
                raise AcceptanceFailure("Wheel must contain exactly one distribution metadata file")
            metadata = BytesParser().parsebytes(archive.read(names[0]))
        kind = "wheel"
        source_contents = None
    elif path.name.endswith(".tar.gz"):
        with tarfile.open(path, "r:gz") as archive:
            members = [item for item in archive.getmembers() if item.isfile()]
            candidates = [item for item in members if item.name.endswith("/PKG-INFO")]
            candidates.sort(key=lambda item: item.name.count("/"))
            if not candidates:
                raise AcceptanceFailure("Sdist is missing PKG-INFO")
            metadata_file = archive.extractfile(candidates[0])
            if metadata_file is None:
                raise AcceptanceFailure("Cannot read source metadata")
            metadata = BytesParser().parsebytes(metadata_file.read())
            root = candidates[0].name.rsplit("/", 1)[0] + "/"
            names = {item.name[len(root):] for item in members if item.name.startswith(root)}
            required = {"README.md", "LICENSE", "pyproject.toml", "examples/boxes.csv",
                        "examples/reviewers.csv", "tests/test_core.py", "tests/test_workflow.py",
                        "scripts/verify_release.py"}
            if not required <= names or not any(name.startswith("docs/") and name.endswith(".md") for name in names):
                raise AcceptanceFailure("Sdist must include documentation, examples, tests and this verifier")
            pyproject_file = archive.extractfile(root + "pyproject.toml")
            if pyproject_file is None:
                raise AcceptanceFailure("Cannot read source build contract")
            pyproject = pyproject_file.read().decode("utf-8")
            section = re.search(r"(?ms)^\[build-system\]\s*\n(.*?)(?=^\[|\Z)", pyproject)
            requires = re.search(r"(?ms)^requires\s*=\s*(\[.*?\])", section.group(1) if section else "")
            if not requires:
                raise AcceptanceFailure("Unsupported source build requirements format")
            # BoxReply uses the Python/TOML common string-list subset. No TOML
            # dependency is needed on Python 3.10, and literal_eval executes no code.
            build_requires = ast.literal_eval(requires.group(1))
            if not isinstance(build_requires, list) or not all(isinstance(v, str) for v in build_requires):
                raise AcceptanceFailure("Source build requirements must be a string list")
            source_contents = {"required_files_present": sorted(required), "documentation_present": True}
        kind = "sdist"
    else:
        raise AcceptanceFailure("Artifact must be a .whl or .tar.gz file")
    if metadata.get("Name", "").lower().replace("_", "-") != "boxreply":
        raise AcceptanceFailure("This acceptance harness supports BoxReply artifacts only")
    runtime_requires = [value for value in metadata.get_all("Requires-Dist", [])
                        if not re.search(r";.*\bextra\b", value)]
    return {"kind": kind, "package": metadata["Name"], "version": metadata["Version"],
            "runtime_requires": runtime_requires, "build_requires": build_requires,
            "source_contents": source_contents}


class Harness:
    def __init__(self, artifact, work, report):
        self.artifact = artifact
        self.work = work
        self.report = report
        self.environment = os.environ.copy()
        for key in ("PYTHONPATH", "PYTHONHOME", "VIRTUAL_ENV"):
            self.environment.pop(key, None)
        self.environment["PYTHONNOUSERSITE"] = "1"
        self.environment["PIP_DISABLE_PIP_VERSION_CHECK"] = "1"
        self.logs = work / "logs"
        self.logs.mkdir()
        self.python = work / "venv" / ("Scripts/python.exe" if os.name == "nt" else "bin/python")

    def sanitized(self, text):
        replacements = [(str(self.work), "$WORK"), (str(self.artifact.parent), "$ARTIFACT_DIR"),
                        (str(Path(__file__).resolve().parent.parent), "$CHECKOUT"), (str(Path.home()), "$HOME")]
        for original, substitute in sorted(replacements, key=lambda pair: len(pair[0]), reverse=True):
            text = text.replace(original, substitute)
        return text

    def run(self, label, args, *, expected=0, cwd=None, timeout=180):
        completed = subprocess.run(list(map(str, args)), cwd=cwd or self.work,
                                   env=self.environment, text=True, capture_output=True, timeout=timeout)
        log = self.logs / f"{len(self.report['checks']) + 1:02d}-{label}.txt"
        log.write_text(self.sanitized(completed.stdout + completed.stderr), encoding="utf-8")
        check = {"name": label, "expected_exit": expected, "observed_exit": completed.returncode,
                 "passed": completed.returncode == expected, "diagnostic": log.relative_to(self.work).as_posix()}
        self.report["checks"].append(check)
        if not check["passed"]:
            raise AcceptanceFailure(f"Check {label} returned {completed.returncode}; expected {expected}")
        if "Traceback" in completed.stderr and expected in {2, 3}:
            raise AcceptanceFailure(f"Check {label} emitted an unexpected traceback")
        return completed

    def cli(self, label, *args, expected=0, cwd=None):
        return self.run(label, [self.python, "-I", "-c", CLI_BOOTSTRAP, *args], expected=expected, cwd=cwd)

    def assert_that(self, label, condition):
        self.report["checks"].append({"name": label, "passed": bool(condition)})
        if not condition:
            raise AcceptanceFailure(f"Acceptance assertion failed: {label}")

    def workbook(self, label, source, target, mode="answer"):
        self.run(label, [self.python, "-I", "-c", EDIT_WORKBOOK, source, target, mode])


def write_csv(path, header, rows):
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(header)
        writer.writerows(rows)


def load_results(packet):
    return json.loads((packet / "results.json").read_text(encoding="utf-8"))


def rehash_receipt(packet):
    path = packet / "receipt.json"
    receipt = json.loads(path.read_text(encoding="utf-8"))
    receipt["files"] = {item.relative_to(packet).as_posix(): digest(item)
                        for item in packet.rglob("*") if item.is_file() and item != path}
    path.write_text(json.dumps(receipt, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def acceptance(harness):
    workspace = harness.work / "acceptance"
    workspace.mkdir()
    inputs = workspace / "inputs"
    inputs.mkdir()
    boxes, reviewers = inputs / "boxes.csv", inputs / "reviewers.csv"
    write_csv(boxes, ("box_id", "department", "record_series", "description", "proposed_action", "proposed_date"), [
        ("000003", "Planning", "MAP-1", "Historic plan drawings", "transfer", "2026-12-31"),
        ("000001", "Finance", "FIN-1", "Invoices, grants and receipts", "destroy", "2026-11-30"),
        ("000002", "Finance", "FIN-2", "Audit working papers", "destroy", "2026-11-30"),
    ])
    write_csv(reviewers, ("department", "reviewer_id", "reviewer_name"), [
        ("Finance", "alex", "Alex Reviewer"), ("Finance", "casey", "Casey Reviewer"),
        ("Planning", "casey", "Casey Reviewer"), ("Planning", "sam", "Sam Reviewer"),
    ])
    review_round = workspace / "round"
    harness.cli("offline-prepare", "prepare", "--boxes", boxes, "--reviewers", reviewers,
                "--round-id", "release-acceptance", "--out", review_round)
    returns = workspace / "returns"
    returns.mkdir()
    answered = []
    for reviewer in ("alex", "casey", "sam"):
        source, = review_round.glob(f"reviews/*-{reviewer}.xlsx")
        target = returns / f"{reviewer}.xlsx"
        harness.workbook(f"edit-generated-{reviewer}", source, target)
        answered.append(target)
    happy = workspace / "happy-packet"
    harness.cli("offline-complete-reconciliation", "reconcile", "--round", review_round,
                "--returns", *answered, "--out", happy)
    observed = load_results(happy)
    harness.assert_that("six-explicit-decisions-and-preserved-identifiers",
                        observed["complete"] and observed["ready_for_coordinator_review"]
                        and observed["counts"]["answered_decisions"] == 6
                        and all(row["reported_no_exception"] for row in observed["rows"])
                        and [row["box_id"] for row in observed["rows"]] == ["000001", "000002", "000003"])
    time.sleep(1.1)  # Prevent same-second workbook timestamps from hiding replay defects.
    harness.cli("offline-verify-after-clock-boundary", "verify", "--packet", happy)
    partial = workspace / "partial-packet"
    harness.cli("offline-partial-return-exit3", "reconcile", "--round", review_round,
                "--returns", answered[0], "--out", partial, expected=3)
    observed = load_results(partial)
    harness.assert_that("partial-never-becomes-no-exception",
                        not observed["complete"] and not observed["ready_for_coordinator_review"]
                        and observed["counts"]["answered_decisions"] == 2
                        and all(not row["reported_no_exception"] for row in observed["rows"]))
    harness.cli("offline-verify-valid-progress-packet", "verify", "--packet", partial)
    exception = returns / "alex-exception.xlsx"
    harness.workbook("edit-exception", answered[0], exception, "exception")
    exception_packet = workspace / "exception-packet"
    harness.cli("offline-exception-with-missing-reviewer", "reconcile", "--round", review_round,
                "--returns", exception, "--out", exception_packet, expected=3)
    observed = load_results(exception_packet)["rows"][0]
    harness.assert_that("exception-and-pending-both-visible",
                        observed["status"] == "exception" and observed["pending_reviewers"] == ["casey"])
    correction = returns / "casey-correction.xlsx"
    harness.workbook("edit-correction", answered[1], correction, "needs_correction")
    correction_packet = workspace / "correction-packet"
    harness.cli("offline-correction-blocks-readiness", "reconcile", "--round", review_round,
                "--returns", answered[0], correction, answered[2], "--out", correction_packet, expected=3)
    observed = load_results(correction_packet)
    harness.assert_that("complete-correction-remains-blocked", observed["complete"] and not observed["ready_for_coordinator_review"])
    changed = returns / "alex-changed.xlsx"
    harness.workbook("edit-wrong-source", answered[0], changed, "wrong_source")
    rejected = workspace / "rejected-source"
    harness.cli("offline-wrong-source-exit2", "reconcile", "--round", review_round,
                "--returns", changed, "--out", rejected, expected=2)
    harness.assert_that("wrong-source-publishes-no-packet", not rejected.exists())
    duplicate = returns / "alex-duplicate.xlsx"
    shutil.copyfile(answered[0], duplicate)
    rejected = workspace / "rejected-duplicate"
    harness.cli("offline-duplicate-return-exit2", "reconcile", "--round", review_round,
                "--returns", answered[0], duplicate, "--out", rejected, expected=2)
    harness.assert_that("duplicate-publishes-no-packet", not rejected.exists())
    relocated = workspace / "relocated packet"
    shutil.move(str(happy), str(relocated))
    shutil.rmtree(review_round)
    shutil.rmtree(returns)
    shutil.rmtree(inputs)
    isolated = workspace / "unrelated working directory"
    isolated.mkdir()
    harness.cli("offline-relocated-source-removed-replay", "verify", "--packet", relocated, cwd=isolated)
    harness.assert_that("original-round-inputs-and-returns-removed",
                        not review_round.exists() and not returns.exists() and not inputs.exists())
    tampered = workspace / "tampered-hash-packet"
    shutil.copytree(relocated, tampered)
    with (tampered / "results.json").open("ab") as stream:
        stream.write(b"changed")
    harness.cli("offline-tampered-hash-rejected", "verify", "--packet", tampered, expected=2)
    rewritten = workspace / "rewritten-receipt-packet"
    shutil.copytree(relocated, rewritten)
    data = load_results(rewritten)
    data["ready_for_coordinator_review"] = False
    (rewritten / "results.json").write_text(json.dumps(data, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    rehash_receipt(rewritten)
    harness.cli("offline-rewritten-receipt-replay-mismatch-rejected", "verify", "--packet", rewritten, expected=2)
    harness.assert_that("all-output-transactions-cleaned", not list(workspace.glob("*.boxreply-lock")) and not list(workspace.glob(".boxreply-*")))
    harness.report["runtime_network_policy"] = "Python sockets and DNS resolution disabled in every application subprocess"
    harness.report["accepted_workflow"] = {"boxes": 3, "required_reviewers": 3, "required_decisions": 6,
                                           "original_source_removed": True, "relocated_packet_replayed": True}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--artifact", type=Path, required=True, help="Downloaded .whl or .tar.gz artifact")
    parser.add_argument("--work", type=Path, required=True, help="New directory for the fresh environment and report")
    args = parser.parse_args(argv)
    if args.artifact.is_symlink() or not args.artifact.is_file():
        parser.error("--artifact must be a regular file")
    if args.work.exists() or args.work.is_symlink():
        parser.error("--work must be a new directory")
    artifact, work = args.artifact.resolve(), args.work.absolute()
    if work.resolve().is_relative_to(Path(__file__).resolve().parent.parent):
        parser.error("--work must be outside the development checkout")
    work.mkdir(parents=True)
    report = {"schema": 1, "product": "BoxReply release acceptance", "status": "running",
              "started_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
              "artifact": {"filename": artifact.name, "sha256": digest(artifact), "bytes": artifact.stat().st_size},
              "harness_sha256": digest(Path(__file__)),
              "host": {"system": platform.system(), "machine": platform.machine(),
                       "python": platform.python_version(), "implementation": platform.python_implementation()},
              "checks": []}
    harness = Harness(artifact, work, report)
    try:
        contract = artifact_contract(artifact)
        report["artifact"].update(contract)
        venv.EnvBuilder(with_pip=True, system_site_packages=False).create(work / "venv")
        harness.assert_that("fresh-venv-excludes-system-site-packages",
                            "include-system-site-packages = false" in (work / "venv/pyvenv.cfg").read_text(encoding="utf-8"))
        wheelhouse = work / "wheelhouse"
        wheelhouse.mkdir()
        requires = list(dict.fromkeys(contract["runtime_requires"] + contract["build_requires"]))
        if requires:
            harness.run("download-runtime-and-build-wheels", [harness.python, "-I", "-m", "pip", "--disable-pip-version-check",
                        "download", "--no-cache-dir", "--only-binary=:all:", "--dest", wheelhouse, *requires], timeout=300)
        candidate = wheelhouse / artifact.name
        shutil.copyfile(artifact, candidate)
        harness.assert_that("artifact-copy-hash-matches-download", digest(candidate) == report["artifact"]["sha256"])
        report["wheelhouse"] = [{"filename": p.name, "sha256": digest(p), "bytes": p.stat().st_size}
                                for p in sorted(wheelhouse.iterdir()) if p.is_file()]
        harness.run("install-artifact-no-index", [harness.python, "-I", "-m", "pip", "--disable-pip-version-check",
                    "install", "--no-cache-dir", "--no-index", "--find-links", wheelhouse, candidate], timeout=300)
        report["installation"] = json.loads(harness.run("installed-origin-and-provenance",
            [harness.python, "-I", "-c", IMPORT_PROBE, Path(__file__).resolve().parent.parent,
             report["artifact"]["sha256"], artifact.name]).stdout)
        harness.assert_that("installed-version-matches-artifact", report["installation"]["version"] == contract["version"])
        console = work / "venv" / ("Scripts/boxreply.exe" if os.name == "nt" else "bin/boxreply")
        process = harness.run("installed-console-entrypoint", [console, "--version"])
        harness.assert_that("console-version-matches-artifact", process.stdout.strip() == contract["version"])
        acceptance(harness)
        report["status"] = "passed"
        report["install_network_policy"] = "Dependencies acquired first; artifact installed with --no-index --find-links and no cache"
    except Exception as exc:
        report["status"] = "failed"
        report["failure"] = {"type": type(exc).__name__, "message": harness.sanitized(str(exc))}
    report["finished_utc"] = dt.datetime.now(dt.timezone.utc).isoformat()
    report["checks_passed"] = sum(bool(check["passed"]) for check in report["checks"])
    serialized = json.dumps(report, indent=2, sort_keys=True) + "\n"
    (work / "report.json").write_text(serialized, encoding="utf-8")
    print(serialized, end="")
    return 0 if report["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
