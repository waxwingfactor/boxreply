"""Small, formula-free workbook interface. Names are declared, not authenticated."""
from __future__ import annotations

import datetime as dt
import io
import re
import zipfile
from xml.etree import ElementTree as StandardXML
from pathlib import Path

from defusedxml import ElementTree
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

from .core import BOX_FIELDS, BoxReplyError, expected_boxes

MAX_XLSX = 10 * 1024 * 1024
MAX_EXPANDED = 50 * 1024 * 1024
RETURN_FIELDS = (*BOX_FIELDS, "decision", "reason")
NAVY = "183C4A"
BLUE = "E7F2FF"
FIXED_TIME = dt.datetime(2000, 1, 1)


def _literal(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # Never interpret spreadsheet formula prefixes.
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _base_sheet(ws, headers, widths):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"


def _append(ws, values, row_number):
    for column, value in enumerate(values, 1):
        cell = ws.cell(row_number, column)
        _literal(cell, value)
        cell.font = Font(name="Calibri", size=11)
    ws.row_dimensions[row_number].height = 48


def _save(wb) -> bytes:
    wb.properties.creator = "BoxReply"
    wb.properties.lastModifiedBy = "BoxReply"
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME
    raw = io.BytesIO()
    wb.save(raw)
    # Canonical archive timestamps make independent report replay byte-exact.
    out = io.BytesIO()
    with zipfile.ZipFile(raw) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in sorted(src.namelist()):
            info = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            content = src.read(name)
            if name == "docProps/core.xml":
                # openpyxl rewrites modified during save, after properties are set.
                content = re.sub(rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)",
                                 rb"\g<1>2000-01-01T00:00:00Z\g<2>", content)
            if name.endswith((".xml", ".rels")):
                content = StandardXML.canonicalize(content.decode("utf-8")).encode("utf-8")
            dst.writestr(info, content)
    wb.close()
    return out.getvalue()


def review_bytes(manifest: dict, reviewer: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 54
    ws["A1"] = "Box review return"
    ws["A1"].font = Font(size=22, bold=True, color=NAVY)
    ws.merge_cells("A1:C1")
    ws["A2"] = "Report exceptions to the human-supplied proposal. This workbook does not authorize disposition."
    ws.merge_cells("A2:C3")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    for row, label, value in [
        (4, "Designated reviewer", reviewer["reviewer_name"]),
        (5, "Submitted by", ""),
        (6, "Submitted on (YYYY-MM-DD)", ""),
        (8, "Round", manifest["round_id"]),
        (10, "1. Review", "Read each original box row on Review. Sort only complete rows."),
        (11, "2. Decide", "Fill decision: no_exception, exception, or needs_correction. Blank remains unanswered."),
        (12, "3. Explain", "For exception or needs_correction, fill reason. Do not edit source columns A:F or remove rows."),
        (13, "4. Return", "Enter your declared name and date above. Save as .xlsx and return through your approved channel."),
        (15, "Identity", "A typed name is a declaration, not an authenticated signature. Keep required authorization evidence separately."),
    ]:
        _literal(ws.cell(row, 1), label)
        ws.cell(row, 1).font = Font(bold=True, color=NAVY, size=11)
        _literal(ws.cell(row, 2), value)
        ws.row_dimensions[row].height = 46 if row >= 10 else 28
    for key in ("B5", "B6"):
        ws[key].fill = PatternFill("solid", fgColor=BLUE)
        ws[key].number_format = "@"
    ws.freeze_panes = "B5"
    ws.print_area = "A1:C15"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    data = wb.create_sheet("Review")
    _base_sheet(data, RETURN_FIELDS, {"A":18,"B":23,"C":21,"D":58,"E":19,"F":20,"G":25,"H":60})
    for row_number, box in enumerate(expected_boxes(manifest, reviewer["reviewer_id"]), 2):
        _append(data, [box[k] for k in BOX_FIELDS] + ["", ""], row_number)
        for column in range(1, 7):
            cell = data.cell(row_number, column)
            cell.number_format = "@"
        for column in (7, 8):
            cell = data.cell(row_number, column)
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.protection = Protection(locked=False)
    data.auto_filter.ref = data.dimensions
    dv = DataValidation(type="list", formula1='"no_exception,exception,needs_correction"', allow_blank=True)
    dv.errorTitle = "Use a supported decision"
    dv.error = "Choose no_exception, exception or needs_correction."
    dv.showErrorMessage = True
    dv.errorStyle = "stop"
    data.add_data_validation(dv)
    dv.add(f"G2:G{data.max_row}")
    meta = wb.create_sheet("_meta")
    for key, value in [("schema", "1"), ("round_id", manifest["round_id"]),
                       ("manifest_sha256", manifest["manifest_sha256"]),
                       ("reviewer_id", reviewer["reviewer_id"]), ("reviewer_name", reviewer["reviewer_name"])]:
        meta.append([key, value])
        _literal(meta.cell(meta.max_row, 1), key)
        _literal(meta.cell(meta.max_row, 2), value)
    meta.sheet_state = "hidden"
    return _save(wb)


def _text(cell, allow_date=False):
    value = cell.value
    if cell.data_type == "f":
        raise BoxReplyError("Formulas are unsupported in return workbooks")
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if allow_date and isinstance(value, (dt.datetime, dt.date)):
        if isinstance(value, dt.datetime) and value.time() != dt.time():
            raise BoxReplyError("Submitted on must contain a date without a time")
        return value.strftime("%Y-%m-%d")
    raise BoxReplyError(f"Cell {cell.coordinate} must be text; preserve identifiers as text")


def read_response(path: Path) -> dict:
    try:
        if path.is_symlink() or not path.is_file() or path.stat().st_size > MAX_XLSX:
            raise BoxReplyError("Return must be a regular .xlsx file of at most 10 MiB")
        raw = path.read_bytes()
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            infos = z.infolist()
            names = [i.filename for i in infos]
            if len(infos) > 200 or len(set(names)) != len(names) or sum(i.file_size for i in infos) > MAX_EXPANDED:
                raise BoxReplyError("Return workbook exceeds archive limits or has duplicate members")
            if any(i.flag_bits & 1 or i.file_size > MAX_EXPANDED for i in infos):
                raise BoxReplyError("Encrypted or oversized workbook member")
            if any("externalLink" in n or "vbaProject" in n or n.endswith(".bin") for n in names):
                raise BoxReplyError("Macros, binary content and external workbook links are unsupported")
            # Relationships may point sheet XML anywhere in the archive. Inspect
            # every XML member, independent of filename or stale dimensions.
            for name in names:
                if not name.endswith((".xml", ".rels")):
                    raise BoxReplyError("Only XML workbook content is supported")
                if name.endswith((".xml", ".rels")):
                    cells = 0
                    with z.open(name) as xml:
                        for _, node in ElementTree.iterparse(xml, events=("end",)):
                            tag = node.tag
                            prefix = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                            if tag in {prefix + name for name in ("externalLink", "externalReferences", "externalReference", "externalBook")}:
                                raise BoxReplyError("External workbook links are unsupported")
                            relationship = node.attrib.get("Type", "").rsplit("/", 1)[-1]
                            if relationship in {"externalLink", "externalLinkPath", "vbaProject", "oleObject", "package"}:
                                raise BoxReplyError("External workbook links, macros and embedded objects are unsupported")
                            content_type = node.attrib.get("ContentType", "").lower()
                            if any(token in content_type for token in ("macroenabled", "vbaproject", "macrosheet", "oleobject")):
                                raise BoxReplyError("Macro or embedded workbook content is unsupported")
                            if tag == prefix + "f":
                                raise BoxReplyError("Formulas are unsupported in return workbooks")
                            if tag == prefix + "row" and not 1 <= int(node.attrib.get("r", "0")) <= 10001:
                                raise BoxReplyError("Return row exceeds supported dimensions")
                            if tag == prefix + "c":
                                cells += 1
                                ref = node.attrib.get("r", "")
                                col = ref.rstrip("0123456789")
                                row = ref[len(col):]
                                if cells > 80008 or col not in "ABCDEFGH" or len(col) != 1 or not row.isascii() or not row.isdigit() or not 1 <= int(row) <= 10001:
                                    raise BoxReplyError("Return cells exceed supported dimensions")
                            node.clear()
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            if set(wb.sheetnames) != {"Instructions", "Review", "_meta"}:
                raise BoxReplyError("Return must keep exactly the generated sheets")
            for sheet in wb:
                if (sheet.max_row is not None and sheet.max_row > 10001) or (sheet.max_column is not None and sheet.max_column > 8):
                    raise BoxReplyError("Return worksheet exceeds supported dimensions")
                # Recalculate dimensions for editors that write stale dimension attributes.
                sheet.reset_dimensions()
            metadata_rows = list(wb["_meta"].iter_rows())
            if len(metadata_rows) != 5 or any(len(row) != 2 for row in metadata_rows):
                raise BoxReplyError("Invalid workbook metadata")
            metadata = {}
            for row in metadata_rows:
                k, v = map(_text, row)
                if k in metadata:
                    raise BoxReplyError("Duplicate workbook metadata")
                metadata[k] = v
            if set(metadata) != {"schema","round_id","manifest_sha256","reviewer_id","reviewer_name"} or metadata.pop("schema") != "1":
                raise BoxReplyError("Unsupported workbook metadata")
            inst = wb["Instructions"]
            if _text(inst["A5"]) != "Submitted by" or _text(inst["A6"]) != "Submitted on (YYYY-MM-DD)":
                raise BoxReplyError("Submission fields have been moved or changed")
            response = {**metadata, "submitted_by": _text(inst["B5"]), "submitted_on": _text(inst["B6"], True), "rows": []}
            ws = wb["Review"]
            rows = ws.iter_rows()
            header = next(rows, ())
            if tuple(_text(c) for c in header) != RETURN_FIELDS:
                raise BoxReplyError("Review columns have changed")
            for row in rows:
                if len(response["rows"]) >= 10000:
                    raise BoxReplyError("Return has too many rows")
                vals = [_text(c, allow_date=(i == 5)) for i, c in enumerate(row)]
                # Some editors retain formatting on trailing empty rows.
                if not any(vals):
                    continue
                if len(vals) != len(RETURN_FIELDS):
                    raise BoxReplyError("Review row has an unsupported column count")
                response["rows"].append(dict(zip(RETURN_FIELDS, vals)))
            return response
        finally:
            wb.close()
    except BoxReplyError:
        raise
    except Exception as e:
        raise BoxReplyError(f"Cannot read return workbook: {type(e).__name__}") from e


def coordinator_bytes(result: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 37
    ws.column_dimensions["B"].width = 82
    ws.sheet_view.showGridLines = False
    summary = [
        ("BoxReply packet", result["round_id"]),
        ("Scope", "Reported exceptions to a frozen, human-supplied box list"),
        ("All required decisions answered", "Yes" if result["complete"] else "No"),
        ("Ready for coordinator review", "Yes" if result["ready_for_coordinator_review"] else "No"),
        ("Next step", "Review every exception and correction, obtain ordinary authorization, and check whether the source list is still current."),
        ("Meaning", "No exception reported is a reviewer assertion. It does not authorize destruction or transfer."),
        ("Round fingerprint", result["manifest_sha256"]),
    ] + list(result["counts"].items())
    for row_number, values in enumerate(summary, 1):
        _append(ws, values, row_number)
        ws.row_dimensions[row_number].height = 26
    ws.row_dimensions[5].height = 42
    ws.row_dimensions[6].height = 42
    ws["A1"].font = Font(size=19, bold=True, color=NAVY)
    ws.row_dimensions[1].height = 35
    ws.freeze_panes = "B3"
    ws.print_area = ws.dimensions
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    boxes = wb.create_sheet("Boxes")
    headers = (*BOX_FIELDS, "status", "pending_reviewers", "correction_reviewers")
    _base_sheet(boxes, headers, {"A":18,"B":22,"C":21,"D":58,"E":19,"F":20,"G":25,"H":32,"I":32})
    decisions = wb.create_sheet("Decisions")
    dfields = ("reviewer_id", "reviewer_name", "returned", "submitted_by", "submitted_on", "decision", "reason")
    _base_sheet(decisions, ("box_id", *dfields), {"A":18,"B":20,"C":26,"D":14,"E":27,"F":20,"G":25,"H":60})
    colors = {"exception":"FCE4D6", "needs_correction":"FFF2CC", "incomplete":"E7E6E6"}
    decision_row = 2
    for row_number, row in enumerate(result["rows"], 2):
        _append(boxes, [row[k] for k in BOX_FIELDS] + [row["status"], ", ".join(row["pending_reviewers"]), ", ".join(row["correction_reviewers"])], row_number)
        if row["status"] in colors:
            boxes.cell(row_number, 7).fill = PatternFill("solid", fgColor=colors[row["status"]])
        for decision in row["decisions"]:
            _append(decisions, [row["box_id"]] + [decision[k] for k in dfields], decision_row)
            decision_row += 1
    boxes.auto_filter.ref = boxes.dimensions
    decisions.auto_filter.ref = decisions.dimensions
    return _save(wb)
