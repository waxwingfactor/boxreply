"""Independent regressions for archive-part portability and literal cells."""
from __future__ import annotations

import csv
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import zipfile

import openpyxl
import pytest

from boxreply.core import BOX_FIELDS, BoxReplyError, prepare, reconcile
from boxreply.xlsx import coordinator_bytes, read_response, review_bytes


NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


@pytest.fixture
def sample(tmp_path):
    for filename, fields, values in [
        ("boxes.csv", BOX_FIELDS,
         [("00001", "Finance", "+Literal series", "=Literal description", "destroy", "2026-10-15")]),
        ("reviewers.csv", ("department", "reviewer_id", "reviewer_name"),
         [("Finance", "reviewer", "=Example")]),
    ]:
        with (tmp_path / filename).open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(fields)
            writer.writerows(values)
    manifest = prepare(tmp_path / "boxes.csv", tmp_path / "reviewers.csv", "portable-2026")
    path = tmp_path / "return.xlsx"
    path.write_bytes(review_bytes(manifest, manifest["reviewers"][0]))
    return manifest, path


def rewrite_archive(path, *, rename=None, transform=None, extra=None):
    rename = rename or {}
    with zipfile.ZipFile(path) as source:
        members = [(info.filename, source.read(info)) for info in source.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for old_name, content in members:
            # OPC relationships and content types must refer to renamed parts.
            for old, new in rename.items():
                content = content.replace(old.encode(), new.encode())
            new_name = rename.get(old_name, old_name)
            if transform:
                content = transform(new_name, content)
            target.writestr(new_name, content)
        for name, content in (extra or {}).items():
            target.writestr(name, content)


def complete_return(path):
    workbook = openpyxl.load_workbook(path)
    for sheet, coordinate, value in [
        ("Instructions", "B5", "=Example"),
        ("Instructions", "B6", "2026-10-15"),
        ("Review", "G2", "no_exception"),
    ]:
        cell = workbook[sheet][coordinate]
        cell.value = value
        cell.data_type = "s"
    workbook.save(path)
    workbook.close()


@pytest.mark.parametrize("alternative", ["xl/review-data.xml", "xl/worksheets/renamed-return.xml"])
def test_valid_alternate_worksheet_part_names_preserve_complete_response(sample, alternative):
    manifest, path = sample
    complete_return(path)
    before = read_response(path)
    rewrite_archive(path, rename={"xl/worksheets/sheet2.xml": alternative})
    observed = read_response(path)
    assert observed == before
    assert reconcile(manifest, [observed])["ready_for_coordinator_review"]


@pytest.mark.parametrize("part", ["xl/worksheets/sheet1.xml", "xl/review-data.xml", "xl/worksheets/sheet3.xml"])
@pytest.mark.parametrize("cell", ["A10002", "I2", "XFD1048576"])
def test_out_of_range_cells_rejected_in_every_referenced_part(sample, part, cell):
    _, path = sample
    rename = {"xl/worksheets/sheet2.xml": "xl/review-data.xml"}

    def transform(name, content):
        if name != part:
            return content
        # A stale dimension deliberately remains valid. Physical cells must be checked.
        injection = f'<row r="2"><c r="{cell}" t="inlineStr"><is><t>unsupported</t></is></c></row>'
        return content.replace(b"</sheetData>", injection.encode() + b"</sheetData>")

    rewrite_archive(path, rename=rename, transform=transform)
    with pytest.raises(BoxReplyError, match="dimensions"):
        read_response(path)


def test_out_of_range_cells_in_unreferenced_xml_cannot_bypass_scan(sample):
    _, path = sample
    xml = f'<worksheet xmlns="{NS}"><sheetData><row r="10002"><c r="A10002" /></row></sheetData></worksheet>'
    rewrite_archive(path, extra={"custom/unused-review.xml": xml.encode()})
    with pytest.raises(BoxReplyError, match="dimensions"):
        read_response(path)


def test_formula_like_reviewer_name_is_literal_in_template_return_and_report(sample):
    manifest, path = sample
    workbook = openpyxl.load_workbook(path, data_only=False)
    for sheet, coordinate in [("Instructions", "B4"), ("_meta", "B5")]:
        assert workbook[sheet][coordinate].value == "=Example"
        assert workbook[sheet][coordinate].data_type == "s"
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()
    assert read_response(path)["reviewer_name"] == "=Example"
    complete_return(path)
    response = read_response(path)
    assert response["submitted_by"] == "=Example"
    result = reconcile(manifest, [response])
    assert result["rows"][0]["reported_no_exception"]
    workbook = openpyxl.load_workbook(io.BytesIO(coordinator_bytes(result)), data_only=False)
    assert workbook["Decisions"]["C2"].value == "=Example"
    assert workbook["Decisions"]["C2"].data_type == "s"
    assert workbook["Decisions"]["E2"].data_type == "s"
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()


@pytest.mark.parametrize("part", ["xl/worksheets/sheet1.xml", "xl/review-data.xml", "xl/worksheets/sheet3.xml", "unreferenced"])
def test_formula_anywhere_in_archive_is_unsupported(sample, part):
    _, path = sample
    rename = {"xl/worksheets/sheet2.xml": "xl/review-data.xml"}

    def transform(name, content):
        if name == part:
            formula = b'<row r="2"><c r="H2"><f>1+1</f><v>2</v></c></row>'
            return content.replace(b"</sheetData>", formula + b"</sheetData>")
        return content

    extra = {}
    if part == "unreferenced":
        extra["custom/formula.xml"] = f'<worksheet xmlns="{NS}"><sheetData><row r="1"><c r="A1"><f>1+1</f></c></row></sheetData></worksheet>'.encode()
    rewrite_archive(path, rename=rename, transform=transform, extra=extra)
    with pytest.raises(BoxReplyError, match="Formulas"):
        read_response(path)


def test_external_workbook_link_with_alternate_part_name_is_unsupported(sample):
    """OPC types identify links; their ZIP filenames are not authoritative."""
    _, path = sample
    office_rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel = "http://schemas.openxmlformats.org/package/2006/relationships"

    def transform(name, content):
        if name == "xl/_rels/workbook.xml.rels":
            relationship = f'<Relationship Type="{office_rel}/externalLink" Target="/custom/linked-book.xml" Id="rId99"/>'
            return content.replace(b"</Relationships>", relationship.encode() + b"</Relationships>")
        if name == "xl/workbook.xml":
            reference = f'<externalReferences><externalReference xmlns:r="{office_rel}" r:id="rId99"/></externalReferences>'
            return content.replace(b"</workbook>", reference.encode() + b"</workbook>")
        if name == "[Content_Types].xml":
            override = '<Override PartName="/custom/linked-book.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml"/>'
            return content.replace(b"</Types>", override.encode() + b"</Types>")
        return content

    extra = {
        "custom/linked-book.xml": f'<externalLink xmlns="{NS}" xmlns:r="{office_rel}"><externalBook r:id="rId1"><sheetNames><sheetName val="Sheet1"/></sheetNames><sheetDataSet/></externalBook></externalLink>'.encode(),
        "custom/_rels/linked-book.xml.rels": f'<Relationships xmlns="{package_rel}"><Relationship Id="rId1" Type="{office_rel}/externalLinkPath" Target="file:///does-not-exist.xlsx" TargetMode="External"/></Relationships>'.encode(),
    }
    rewrite_archive(path, transform=transform, extra=extra)
    with pytest.raises(BoxReplyError, match="[Ee]xternal|unsupported"):
        read_response(path)


def test_benign_hyperlink_relationship_remains_literal_report_text(sample):
    manifest, path = sample
    complete_return(path)
    workbook = openpyxl.load_workbook(path)
    workbook["Review"]["H2"] = "Reference in departmental procedure"
    workbook["Review"]["H2"].hyperlink = "https://invalid.example/procedure"
    workbook.save(path)
    workbook.close()
    response = read_response(path)
    result = reconcile(manifest, [response])
    workbook = openpyxl.load_workbook(io.BytesIO(coordinator_bytes(result)))
    assert workbook["Decisions"]["H2"].value == "Reference in departmental procedure"
    assert workbook["Decisions"]["H2"].hyperlink is None
    workbook.close()


def test_native_libreoffice_resave_preserves_review_contract(sample, tmp_path):
    executable = os.environ.get("BOXREPLY_TEST_SOFFICE") or shutil.which("soffice")
    if not executable:
        pytest.skip("LibreOffice is not installed; native resave is an optional local check")
    manifest, path = sample
    complete_return(path)
    before = read_response(path)
    target = tmp_path / "native-resave"
    target.mkdir()
    profile = tmp_path / "libreoffice-profile"
    result = subprocess.run(
        [executable, f"-env:UserInstallation={profile.as_uri()}", "--headless",
         "--convert-to", "xlsx", "--outdir", str(target), str(path)],
        text=True, capture_output=True, timeout=60,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    converted = target / path.name
    assert converted.is_file(), result.stdout + result.stderr
    observed = read_response(converted)
    assert observed == before
    assert reconcile(manifest, [observed])["ready_for_coordinator_review"]


def test_xml_backends_generate_identical_reports_and_replay_packets(sample, tmp_path):
    """lxml is optional in openpyxl and must not change receipt replay bytes.

    BOXREPLY_TEST_LXML_PATH permits a bundled optional lxml installation without
    adding it to the product dependencies. CI environments lacking it skip this
    additional compatibility check.
    """
    from boxreply.cli import main

    manifest, returned_path = sample
    complete_return(returned_path)
    round_path = tmp_path / "round"
    packet_path = tmp_path / "packet"
    assert main(["prepare", "--boxes", str(tmp_path / "boxes.csv"), "--reviewers",
                 str(tmp_path / "reviewers.csv"), "--round-id", manifest["round_id"],
                 "--out", str(round_path)]) == 0
    assert main(["reconcile", "--round", str(round_path), "--returns", str(returned_path),
                 "--out", str(packet_path)]) == 0
    code = """
import hashlib, json, sys
if sys.argv[1]:
    sys.path.append(sys.argv[1])
import openpyxl
from boxreply.cli import verify_packet
from boxreply.xlsx import coordinator_bytes
from pathlib import Path
result = verify_packet(Path(sys.argv[2]))
print(json.dumps({"lxml": openpyxl.LXML, "sha256": hashlib.sha256(coordinator_bytes(result)).hexdigest()}))
"""
    outputs = []
    for enabled in ("True", "False"):
        env = os.environ.copy()
        env["OPENPYXL_LXML"] = enabled
        process = subprocess.run(
            [sys.executable, "-c", code, os.environ.get("BOXREPLY_TEST_LXML_PATH", ""), str(packet_path)],
            env=env, text=True, capture_output=True, timeout=30,
        )
        assert process.returncode == 0, process.stdout + process.stderr
        observed = json.loads(process.stdout)
        if enabled == "True" and not observed["lxml"]:
            pytest.skip("Optional lxml is not installed; stdlib serialization and replay passed")
        assert observed["lxml"] == (enabled == "True")
        outputs.append(observed["sha256"])
    assert outputs[0] == outputs[1]
