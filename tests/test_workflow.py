"""Independent end-to-end and adversarial tests of the published CLI contract.

These tests edit the generated files as a spreadsheet reviewer would, then run
the installed entry point in subprocesses outside the development directory.
"""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import time
import zipfile

import openpyxl
import pytest


BOX_HEADER = ("box_id", "department", "record_series", "description", "proposed_action", "proposed_date")
REVIEWER_HEADER = ("department", "reviewer_id", "reviewer_name")


def command(cwd: Path, *args, expected=None):
    env = os.environ.copy()
    env.pop("PYTHONPATH", None)
    env.pop("PYTHONHOME", None)
    result = subprocess.run(
        [sys.executable, "-m", "boxreply", *map(str, args)],
        cwd=cwd, env=env, text=True, capture_output=True, timeout=30,
    )
    assert "Traceback" not in result.stderr, result.stderr
    if expected is not None:
        assert result.returncode == expected, result.stdout + result.stderr
    return result


def write_csv(path, header, rows):
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(header)
        writer.writerows(rows)


@pytest.fixture
def review_round(tmp_path):
    write_csv(tmp_path / "boxes.csv", BOX_HEADER, [
        ("000003", "Planning", "MAP-1", "Historic plan drawings", "transfer", "2026-12-31"),
        ("000001", "Finance", "FIN-1", "Invoices, grants and receipts", "destroy", "2026-11-30"),
        ("000002", "Finance", "FIN-2", "Audit working papers", "destroy", "2026-11-30"),
    ])
    write_csv(tmp_path / "reviewers.csv", REVIEWER_HEADER, [
        ("Finance", "alex", "Alex Reviewer"),
        ("Finance", "casey", "Casey Reviewer"),
        ("Planning", "casey", "Casey Reviewer"),
        ("Planning", "sam", "Sam Reviewer"),
    ])
    root = tmp_path / "round"
    command(tmp_path, "prepare", "--boxes", tmp_path / "boxes.csv", "--reviewers",
            tmp_path / "reviewers.csv", "--round-id", "autumn-2026", "--out", root, expected=0)
    (tmp_path / "returns").mkdir()
    return root


def returned(review_round, reviewer, *, decisions=None, submitted_by=None, submitted_on="2026-10-15", edit=None):
    source, = review_round.glob(f"reviews/*-{reviewer}.xlsx")
    target = review_round.parent / "returns" / f"{reviewer}.xlsx"
    workbook = openpyxl.load_workbook(source)
    instructions = workbook["Instructions"]
    instructions["B5"] = submitted_by or f"Declared {reviewer}"
    instructions["B6"] = submitted_on
    for row in workbook["Review"].iter_rows(min_row=2):
        decision, reason = (decisions or {}).get(row[0].value, ("no_exception", ""))
        row[6].value = decision
        row[7].value = reason
    if edit:
        edit(workbook)
    workbook.save(target)
    workbook.close()
    return target


def reconcile(review_round, returns=(), *, out=None, expected=None):
    out = out or review_round.parent / "packet"
    args = ["reconcile", "--round", review_round, "--out", out]
    if returns:
        args += ["--returns", *returns]
    return command(review_round.parent, *args, expected=expected), out


def result(packet):
    return json.loads((packet / "results.json").read_text(encoding="utf-8"))


def rejected(review_round, returns=(), *, out=None):
    process, packet = reconcile(review_round, returns, out=out, expected=2)
    assert process.stderr.startswith("boxreply:"), process.stderr
    assert not packet.exists()
    assert not packet.with_name(packet.name + ".boxreply-lock").exists()
    assert not list(packet.parent.glob(".boxreply-*"))
    return process


def change_zip(path, transform):
    with zipfile.ZipFile(path) as original:
        members = [(info, original.read(info)) for info in original.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for info, content in members:
            target.writestr(info, transform(info.filename, content))


def rehash_receipt(packet):
    path = packet / "receipt.json"
    receipt = json.loads(path.read_text(encoding="utf-8"))
    receipt["files"] = {
        p.relative_to(packet).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
        for p in packet.rglob("*") if p.is_file() and p != path
    }
    path.write_text(json.dumps(receipt, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def test_complete_cli_workflow_preserves_original_evidence_and_sorted_rows(review_round):
    files = [returned(review_round, r) for r in ("sam", "alex", "casey")]
    original_bytes = {p.name: p.read_bytes() for p in files}
    _, packet = reconcile(review_round, files, expected=0)
    observed = result(packet)
    assert observed["complete"] and observed["ready_for_coordinator_review"]
    assert observed["counts"] == {
        "boxes": 3, "no_exception": 3, "exception": 0, "needs_correction": 0,
        "incomplete": 0, "required_decisions": 6, "answered_decisions": 6,
        "required_reviewers": 3, "returned_reviewers": 3,
    }
    assert [r["box_id"] for r in observed["rows"]] == ["000001", "000002", "000003"]
    assert all(r["reported_no_exception"] for r in observed["rows"])
    for index, source in enumerate(files, 1):
        assert (packet / "evidence" / "returns" / f"{index:03d}.xlsx").read_bytes() == original_bytes[source.name]
    assert (packet / "evidence/round/boxes.csv").read_bytes() == (review_round / "boxes.csv").read_bytes()
    command(review_round.parent, "verify", "--packet", packet, expected=0)
    workbook = openpyxl.load_workbook(packet / "coordinator.xlsx", data_only=False)
    assert workbook.sheetnames == ["Summary", "Boxes", "Decisions"]
    assert workbook["Boxes"]["A2"].value == "000001"
    assert workbook["Decisions"].max_row == 7
    workbook.close()


def test_packet_replay_survives_a_wall_clock_second_boundary(review_round):
    """ZIP and embedded document properties must both be reproducible."""
    _, packet = reconcile(review_round, expected=3)
    time.sleep(1.1)
    command(review_round.parent, "verify", "--packet", packet, expected=0)


def test_sorting_whole_rows_and_date_cells_is_supported(review_round):
    def reverse_rows(workbook):
        sheet = workbook["Review"]
        values = list(sheet.values)[1:]
        for row_number, values in enumerate(reversed(values), 2):
            for column, value in enumerate(values, 1):
                sheet.cell(row_number, column, value)
        workbook["Instructions"]["B6"] = dt.datetime(2026, 10, 15)
    paths = [returned(review_round, "alex"), returned(review_round, "casey", edit=reverse_rows), returned(review_round, "sam")]
    _, packet = reconcile(review_round, paths, expected=0)
    assert result(packet)["counts"]["answered_decisions"] == 6


def test_missing_return_creates_replayable_progress_packet(review_round):
    _, packet = reconcile(review_round, [returned(review_round, "alex")], expected=3)
    observed = result(packet)
    assert not observed["complete"] and not observed["ready_for_coordinator_review"]
    assert observed["counts"]["answered_decisions"] == 2
    assert observed["rows"][0]["pending_reviewers"] == ["casey"]
    assert observed["rows"][2]["pending_reviewers"] == ["casey", "sam"]
    command(review_round.parent, "verify", "--packet", packet, expected=0)


def test_no_returns_are_not_silent_approval(review_round):
    _, packet = reconcile(review_round, expected=3)
    observed = result(packet)
    assert observed["counts"]["answered_decisions"] == 0
    assert observed["counts"]["incomplete"] == 3
    assert all(not r["reported_no_exception"] for r in observed["rows"])


def test_exception_remains_visible_with_pending_and_blank_reviewer(review_round):
    alex = returned(review_round, "alex", decisions={"000001": ("exception", "Still needed for active review")})
    casey = returned(review_round, "casey", decisions={"000001": ("", "")})
    _, packet = reconcile(review_round, [casey, alex], expected=3)
    first = result(packet)["rows"][0]
    assert first["status"] == "exception" and not first["reported_no_exception"]
    assert first["pending_reviewers"] == ["casey"]
    assert first["decisions"][1]["returned"] is True
    workbook = openpyxl.load_workbook(packet / "coordinator.xlsx")
    assert workbook["Boxes"]["G2"].value == "exception"
    assert workbook["Boxes"]["H2"].value == "casey"
    assert "Still needed" in workbook["Decisions"]["H2"].value
    workbook.close()


def test_complete_exception_is_not_hidden_by_other_no_exception_decisions(review_round):
    paths = [returned(review_round, "alex", decisions={"000002": ("exception", "Audit work continues")}),
             returned(review_round, "casey"), returned(review_round, "sam")]
    _, packet = reconcile(review_round, paths, expected=0)
    observed = result(packet)
    assert observed["complete"] and observed["ready_for_coordinator_review"]
    assert observed["rows"][1]["status"] == "exception"
    assert not observed["rows"][1]["reported_no_exception"]


def test_correction_blocks_readiness_even_under_exception_precedence(review_round):
    paths = [returned(review_round, "alex", decisions={"000001": ("exception", "Retain for current work")}),
             returned(review_round, "casey", decisions={"000001": ("needs_correction", "Description is wrong")}),
             returned(review_round, "sam")]
    _, packet = reconcile(review_round, paths, expected=3)
    observed = result(packet)
    assert observed["complete"] and not observed["ready_for_coordinator_review"]
    assert observed["rows"][0]["status"] == "exception"
    assert observed["rows"][0]["correction_reviewers"] == ["casey"]
    command(review_round.parent, "verify", "--packet", packet, expected=0)


@pytest.mark.parametrize("column,value", [
    (1, "FOREIGN"), (2, "Planning"), (3, "OTHER-SERIES"),
    (4, "Substituted source description"), (5, "transfer"), (6, "2027-11-30"),
])
def test_every_source_column_is_immutable_in_real_return(review_round, column, value):
    path = returned(review_round, "alex", edit=lambda wb: setattr(wb["Review"].cell(2, column), "value", value))
    rejected(review_round, [path])


@pytest.mark.parametrize("mutation", ["deleted", "duplicate", "numeric_id", "new_box"])
def test_missing_duplicate_or_nontext_box_identity_is_rejected(review_round, mutation):
    def edit(workbook):
        sheet = workbook["Review"]
        if mutation == "deleted":
            sheet.delete_rows(2)
        elif mutation == "duplicate":
            sheet.append([c.value for c in sheet[2]])
        elif mutation == "numeric_id":
            sheet["A2"] = 1
        else:
            sheet.append(["FOREIGN", "Finance", "FIN-X", "Injected", "destroy", "2026-11-30", "no_exception", ""])
    rejected(review_round, [returned(review_round, "alex", edit=edit)])


@pytest.mark.parametrize("cell,value", [("B2", "older-round"), ("B3", "0" * 64), ("B4", "sam"), ("B5", "Changed name")])
def test_round_snapshot_and_reviewer_identity_are_bound(review_round, cell, value):
    rejected(review_round, [returned(review_round, "alex", edit=lambda wb: setattr(wb["_meta"][cell], "value", value))])


def test_two_versions_of_one_reviewer_are_rejected_not_last_writer_wins(review_round):
    original = returned(review_round, "alex")
    copy = original.with_name("alex-copy.xlsx")
    shutil.copyfile(original, copy)
    workbook = openpyxl.load_workbook(copy)
    workbook["Review"]["G2"] = "exception"
    workbook["Review"]["H2"] = "Newer opinion"
    workbook.save(copy)
    workbook.close()
    process = rejected(review_round, [original, copy])
    assert "Duplicate submission" in process.stderr


@pytest.mark.parametrize("decision,reason", [("approve", ""), ("exception", ""), ("needs_correction", ""), ("", "Orphan reason")])
def test_invalid_decision_or_missing_required_explanation_is_rejected(review_round, decision, reason):
    rejected(review_round, [returned(review_round, "alex", decisions={"000001": (decision, reason)})])


@pytest.mark.parametrize("sheet,cell", [("Review", "G2"), ("Review", "H2"), ("Instructions", "B5"), ("_meta", "B2")])
def test_formulas_are_never_interpreted_as_decisions_or_metadata(review_round, sheet, cell):
    path = returned(review_round, "alex", edit=lambda wb: setattr(wb[sheet][cell], "value", '=HYPERLINK("https://invalid.example","no_exception")'))
    process = rejected(review_round, [path])
    assert "Formula" in process.stderr


@pytest.mark.parametrize("metadata", ["missing_name", "missing_date", "timestamp", "bad_date"])
def test_decisions_require_declared_name_and_date_without_time(review_round, metadata):
    def edit(workbook):
        if metadata == "missing_name":
            workbook["Instructions"]["B5"] = None
        elif metadata == "missing_date":
            workbook["Instructions"]["B6"] = None
        elif metadata == "timestamp":
            workbook["Instructions"]["B6"] = dt.datetime(2026, 10, 15, 14, 1)
        else:
            workbook["Instructions"]["B6"] = "2026-02-30"
    rejected(review_round, [returned(review_round, "alex", edit=edit)])


@pytest.mark.parametrize("mutation", ["huge_declared_dimension", "sparse_far_row", "far_column", "row_number", "xml_entity", "duplicate_member", "binary", "external_link", "too_many_members", "expanded_limit"])
def test_malicious_xlsx_archives_fail_cleanly_before_publish(review_round, mutation):
    path = returned(review_round, "alex")
    if mutation == "duplicate_member":
        with zipfile.ZipFile(path, "a") as archive:
            with pytest.warns(UserWarning, match="Duplicate name"):
                archive.writestr("xl/workbook.xml", archive.read("xl/workbook.xml"))
    elif mutation in {"binary", "external_link", "too_many_members", "expanded_limit"}:
        with zipfile.ZipFile(path, "a", zipfile.ZIP_DEFLATED) as archive:
            if mutation == "binary":
                archive.writestr("xl/vbaProject.bin", b"unsupported")
            elif mutation == "external_link":
                archive.writestr("xl/externalLinks/externalLink1.xml", b"<externalLink/>")
            elif mutation == "too_many_members":
                for number in range(201):
                    archive.writestr(f"unknown/{number}.xml", b"<x/>")
            else:
                archive.writestr("oversized.xml", b"x" * (50 * 1024 * 1024 + 1))
    else:
        def transform(name, raw):
            if name != "xl/worksheets/sheet2.xml":
                return raw
            if mutation == "huge_declared_dimension":
                return re.sub(br'<dimension ref="[^"]+"', b'<dimension ref="A1:XFD1048576"', raw)
            if mutation == "sparse_far_row":
                return raw.replace(b"</sheetData>", b'<row r="1048576"><c r="A1048576" t="inlineStr"><is><t>x</t></is></c></row></sheetData>')
            if mutation == "far_column":
                return raw.replace(b'r="A2"', b'r="XFD2"')
            if mutation == "row_number":
                return raw.replace(b'<row r="2"', b'<row r="not-an-integer"')
            return b'<!DOCTYPE worksheet [<!ENTITY private SYSTEM "file:///etc/passwd">]>' + raw.replace(b"Invoices, grants and receipts", b"&private;")
        change_zip(path, transform)
    rejected(review_round, [path])


def test_formula_like_source_text_is_literal_in_workbooks(tmp_path):
    write_csv(tmp_path / "boxes.csv", BOX_HEADER, [("000001", "Finance", "+SUM(1,2)", '=HYPERLINK("https://invalid.example","literal")', "destroy", "2026-11-30")])
    write_csv(tmp_path / "reviewers.csv", REVIEWER_HEADER, [("Finance", "alex", "@literal reviewer")])
    root = tmp_path / "round"
    command(tmp_path, "prepare", "--boxes", tmp_path / "boxes.csv", "--reviewers", tmp_path / "reviewers.csv", "--round-id", "literal", "--out", root, expected=0)
    workbook = openpyxl.load_workbook(next(root.glob("reviews/*.xlsx")), data_only=False)
    assert workbook["Review"]["D2"].data_type == "s"
    assert workbook["Review"]["C2"].value == "+SUM(1,2)"
    workbook.close()
    (tmp_path / "returns").mkdir()
    _, packet = reconcile(root, [returned(root, "alex")], expected=0)
    workbook = openpyxl.load_workbook(packet / "coordinator.xlsx", data_only=False)
    assert workbook["Boxes"]["D2"].data_type == "s"
    assert workbook["Boxes"]["D2"].value.startswith("=HYPERLINK")
    assert not any(cell.data_type == "f" for ws in workbook for row in ws for cell in row)
    workbook.close()


def test_existing_output_is_never_replaced(review_round):
    packet = review_round.parent / "existing"
    packet.mkdir()
    marker = packet / "important.txt"
    marker.write_text("keep these exact bytes", encoding="utf-8")
    reconcile(review_round, out=packet, expected=2)
    assert marker.read_text(encoding="utf-8") == "keep these exact bytes"
    assert list(packet.iterdir()) == [marker]


def test_output_inside_round_is_rejected_without_changing_round(review_round):
    before = {p.relative_to(review_round): p.read_bytes() for p in review_round.rglob("*") if p.is_file()}
    reconcile(review_round, out=review_round / "packet", expected=2)
    after = {p.relative_to(review_round): p.read_bytes() for p in review_round.rglob("*") if p.is_file()}
    assert before == after


@pytest.mark.parametrize("where", ["csv_input", "return", "round_file", "round_directory", "output", "packet_evidence"])
def test_symlink_inputs_and_evidence_are_rejected(review_round, where):
    try:
        probe = review_round.parent / "symlink-probe"
        probe.symlink_to(review_round, target_is_directory=True)
        probe.unlink()
    except (OSError, NotImplementedError):
        pytest.skip("Symlink creation is unavailable on this platform")
    if where == "csv_input":
        linked = review_round.parent / "linked.csv"
        linked.symlink_to(review_round / "boxes.csv")
        out = review_round.parent / "new-round"
        command(review_round.parent, "prepare", "--boxes", linked, "--reviewers", review_round / "reviewers.csv", "--round-id", "new", "--out", out, expected=2)
        assert not out.exists()
    elif where == "return":
        path = returned(review_round, "alex")
        linked = path.with_name("linked.xlsx")
        linked.symlink_to(path)
        rejected(review_round, [linked])
    elif where == "round_file":
        path = review_round / "boxes.csv"
        saved = review_round.parent / "saved.csv"
        path.rename(saved)
        path.symlink_to(saved)
        rejected(review_round)
    elif where == "round_directory":
        linked = review_round.parent / "linked-round"
        linked.symlink_to(review_round, target_is_directory=True)
        rejected(linked)
    elif where == "output":
        out = review_round.parent / "linked-output"
        out.symlink_to(review_round.parent / "nonexistent", target_is_directory=True)
        reconcile(review_round, out=out, expected=2)
        assert out.is_symlink()
        assert not out.exists()
    else:
        _, packet = reconcile(review_round, expected=3)
        path = packet / "evidence/round/boxes.csv"
        saved = review_round.parent / "original-evidence.csv"
        path.rename(saved)
        path.symlink_to(saved)
        command(review_round.parent, "verify", "--packet", packet, expected=2)


def test_output_lock_prevents_concurrent_writer_and_is_not_removed(review_round):
    packet = review_round.parent / "packet"
    lock = packet.with_name("packet.boxreply-lock")
    lock.write_text("existing run", encoding="utf-8")
    reconcile(review_round, out=packet, expected=2)
    assert lock.read_text(encoding="utf-8") == "existing run"
    assert not packet.exists()


def test_transaction_rolls_back_if_report_generation_fails(review_round, monkeypatch, capsys):
    from boxreply import cli
    def fail_report(_result):
        raise OSError("Injected full disk during report generation")
    monkeypatch.setattr(cli, "coordinator_bytes", fail_report)
    packet = review_round.parent / "packet"
    code = cli.main(["reconcile", "--round", str(review_round), "--out", str(packet)])
    assert code == 2
    assert "Injected full disk" in capsys.readouterr().err
    assert not packet.exists()
    assert not packet.with_name("packet.boxreply-lock").exists()
    assert not list(packet.parent.glob(".boxreply-*"))


def test_prepare_rolls_back_after_bad_source(review_round):
    broken = review_round.parent / "bad.csv"
    broken.write_text("wrong,columns\n", encoding="utf-8")
    out = review_round.parent / "invalid-round"
    command(review_round.parent, "prepare", "--boxes", broken, "--reviewers", review_round / "reviewers.csv", "--round-id", "broken", "--out", out, expected=2)
    assert not out.exists()
    assert not list(out.parent.glob(".boxreply-*"))
    assert not out.with_name("invalid-round.boxreply-lock").exists()


def test_packet_verifies_after_relocation_without_original_round_or_returns(review_round):
    paths = [returned(review_round, r) for r in ("alex", "casey", "sam")]
    _, packet = reconcile(review_round, paths, expected=0)
    relocated = review_round.parent / "relocated packet"
    shutil.move(packet, relocated)
    shutil.rmtree(review_round)
    shutil.rmtree(review_round.parent / "returns")
    (review_round.parent / "boxes.csv").unlink()
    (review_round.parent / "reviewers.csv").unlink()
    isolated = review_round.parent / "unrelated-cwd"
    isolated.mkdir()
    command(isolated, "verify", "--packet", relocated, expected=0)


@pytest.mark.parametrize("target", ["results.json", "coordinator.xlsx", "evidence/round/boxes.csv", "evidence/returns/001.xlsx"])
def test_changed_packet_files_fail_receipt_hash_check(review_round, target):
    _, packet = reconcile(review_round, [returned(review_round, "alex")], expected=3)
    path = packet / target
    with path.open("ab") as stream:
        stream.write(b"tamper")
    process = command(review_round.parent, "verify", "--packet", packet, expected=2)
    assert "receipt hashes" in process.stderr


@pytest.mark.parametrize("target", ["results", "report", "source", "return", "round_manifest", "unexpected_file"])
def test_replay_detects_semantic_tampering_even_if_receipt_hashes_are_rewritten(review_round, target):
    _, packet = reconcile(review_round, [returned(review_round, "alex")], expected=3)
    if target == "results":
        path = packet / "results.json"
        observed = result(packet)
        observed["ready_for_coordinator_review"] = True
        path.write_text(json.dumps(observed, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    elif target == "report":
        path = packet / "coordinator.xlsx"
        workbook = openpyxl.load_workbook(path)
        workbook["Summary"]["B4"] = True
        workbook.save(path)
        workbook.close()
    elif target == "source":
        path = packet / "evidence/round/boxes.csv"
        path.write_bytes(path.read_bytes().replace(b"Audit working papers", b"Another description"))
    elif target == "return":
        path = packet / "evidence/returns/001.xlsx"
        workbook = openpyxl.load_workbook(path)
        workbook["Review"]["G2"] = "exception"
        workbook["Review"]["H2"] = "Changed after packet generation"
        workbook.save(path)
        workbook.close()
    elif target == "round_manifest":
        path = packet / "evidence/round/manifest.json"
        manifest = json.loads(path.read_text(encoding="utf-8"))
        manifest["boxes"][0]["description"] = "Forged snapshot"
        path.write_text(json.dumps(manifest), encoding="utf-8")
    else:
        (packet / "extra.txt").write_text("untracked addition", encoding="utf-8")
    rehash_receipt(packet)
    command(review_round.parent, "verify", "--packet", packet, expected=2)


@pytest.mark.parametrize("mutation", ["round", "version", "schema_bool", "files_list", "duplicate_key", "missing_receipt"])
def test_malformed_or_inconsistent_receipts_fail_cleanly(review_round, mutation):
    _, packet = reconcile(review_round, expected=3)
    path = packet / "receipt.json"
    receipt = json.loads(path.read_text(encoding="utf-8"))
    if mutation == "round":
        receipt["round_id"] = "different"
    elif mutation == "version":
        receipt["version"] = "999.0"
    elif mutation == "schema_bool":
        receipt["schema"] = True
    elif mutation == "files_list":
        receipt["files"] = []
    elif mutation == "duplicate_key":
        path.write_text('{"schema":1,"schema":1}', encoding="utf-8")
    else:
        path.unlink()
    if mutation not in {"duplicate_key", "missing_receipt"}:
        path.write_text(json.dumps(receipt), encoding="utf-8")
    command(review_round.parent, "verify", "--packet", packet, expected=2)
